"""
Export complet des transactions Firestore vers un fichier Excel.

Usage:
    FIREBASE_CREDENTIALS='<json>' python scripts/export_transactions_xlsx.py
    ou
    python scripts/export_transactions_xlsx.py  # si credentials.json existe à la racine

Options:
    --output <path>   Fichier de sortie (défaut: transactions-export-YYYY-MM-DD.xlsx)
    --since  <YYYY-MM-DD>  Exporter seulement depuis cette date (défaut: tout)
"""
import argparse
import json
import os
import sys
from datetime import date, datetime, timezone
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ─── Firebase init ───────────────────────────────────────────────────────────

def _init_firebase():
    import firebase_admin
    from firebase_admin import credentials, firestore

    cred_json = os.environ.get("FIREBASE_CREDENTIALS")
    if cred_json:
        cred = credentials.Certificate(json.loads(cred_json))
    else:
        cred_path = Path(__file__).parent.parent / "credentials.json"
        if not cred_path.exists():
            sys.exit(
                "❌  Aucun credentials trouvé.\n"
                "    Définir FIREBASE_CREDENTIALS ou placer credentials.json à la racine du projet."
            )
        cred = credentials.Certificate(str(cred_path))

    if not firebase_admin._apps:
        firebase_admin.initialize_app(cred)
    return firestore.client()


# ─── Fetch ───────────────────────────────────────────────────────────────────

COLUMNS = [
    ("ID",              "id"),
    ("Date",            "date"),
    ("Libellé",         "libelle"),
    ("Montant (€)",     "montant"),
    ("Compte",          "compte"),
    ("Catégorie",       "categorie"),
    ("Commentaire",     "commentaire"),
    ("Pointé",          "pointe"),
    ("Mois affectation","moisAffectation"),
    ("Source",          "source"),
    ("En attente",      "enAttente"),
]

COL_WIDTHS = [45, 12, 40, 14, 18, 22, 30, 8, 16, 12, 10]

HEADER_FILL  = PatternFill("solid", fgColor="1A1A2E")
HEADER_FONT  = Font(bold=True, color="D4AF37", size=10)
INCOME_FILL  = PatternFill("solid", fgColor="0D2B0D")
EXPENSE_FILL = PatternFill("solid", fgColor="1E0808")


def fetch_transactions(db, since: str | None = None) -> list[dict]:
    from google.cloud.firestore_v1.base_query import FieldFilter

    query = db.collection("transactions").order_by("date", direction="DESCENDING")
    if since:
        query = query.where(filter=FieldFilter("date", ">=", since))

    print(f"  Chargement des transactions depuis Firestore...", end=" ", flush=True)
    docs = list(query.stream())
    print(f"{len(docs)} documents récupérés.")
    return [{"id": d.id, **(d.to_dict() or {})} for d in docs]


# ─── Excel builder ───────────────────────────────────────────────────────────

def build_workbook(transactions: list[dict]) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transactions"

    # Header
    for col_idx, (header, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = COL_WIDTHS[col_idx - 1]

    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"

    # Data rows
    for row_idx, tx in enumerate(transactions, start=2):
        montant = tx.get("montant")
        is_income = isinstance(montant, (int, float)) and montant > 0
        row_fill = INCOME_FILL if is_income else EXPENSE_FILL

        for col_idx, (_, key) in enumerate(COLUMNS, start=1):
            val = tx.get(key, "")

            if key == "pointe":
                val = "Oui" if val else "Non"
            elif key == "enAttente":
                val = "Oui" if val else ""
            elif key == "montant" and isinstance(val, (int, float)):
                cell = ws.cell(row=row_idx, column=col_idx, value=round(float(val), 2))
                cell.number_format = "#,##0.00"
                cell.fill = row_fill
                continue
            elif val is None:
                val = ""
            else:
                # Firestore Timestamp → str
                val_str = str(val)
                if "DatetimeWithNanoseconds" in type(val).__name__ or hasattr(val, "isoformat"):
                    try:
                        val = val.isoformat()[:10]
                    except Exception:
                        val = val_str

            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = row_fill

    # Auto-filter
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"

    return wb


# ─── Stats sheet ─────────────────────────────────────────────────────────────

def add_stats_sheet(wb: openpyxl.Workbook, transactions: list[dict]) -> None:
    ws = wb.create_sheet("Résumé")

    by_month: dict[str, dict] = {}
    by_cat: dict[str, float]  = {}

    for tx in transactions:
        montant = tx.get("montant")
        if not isinstance(montant, (int, float)):
            continue
        date_str = str(tx.get("date", ""))[:7]  # YYYY-MM
        cat = str(tx.get("categorie") or "Non catégorisé").strip() or "Non catégorisé"

        by_month.setdefault(date_str, {"recettes": 0.0, "depenses": 0.0})
        if montant > 0:
            by_month[date_str]["recettes"] += montant
        else:
            by_month[date_str]["depenses"] += montant

        if montant < 0:
            by_cat[cat] = by_cat.get(cat, 0.0) + montant

    # Header style helper
    def h(ws, row, col, val):
        cell = ws.cell(row=row, column=col, value=val)
        cell.font = Font(bold=True, color="D4AF37")
        cell.fill = PatternFill("solid", fgColor="1A1A2E")

    # Monthly summary
    h(ws, 1, 1, "Mois"); h(ws, 1, 2, "Recettes (€)"); h(ws, 1, 3, "Dépenses (€)"); h(ws, 1, 4, "Solde (€)")
    for i, (month, vals) in enumerate(sorted(by_month.items(), reverse=True), start=2):
        ws.cell(row=i, column=1, value=month)
        ws.cell(row=i, column=2, value=round(vals["recettes"], 2)).number_format = "#,##0.00"
        ws.cell(row=i, column=3, value=round(vals["depenses"], 2)).number_format = "#,##0.00"
        ws.cell(row=i, column=4, value=round(vals["recettes"] + vals["depenses"], 2)).number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 14

    # Category breakdown (col 6+)
    h(ws, 1, 6, "Catégorie"); h(ws, 1, 7, "Total dépenses (€)")
    for i, (cat, total) in enumerate(sorted(by_cat.items(), key=lambda x: x[1]), start=2):
        ws.cell(row=i, column=6, value=cat)
        ws.cell(row=i, column=7, value=round(total, 2)).number_format = "#,##0.00"
    ws.column_dimensions["F"].width = 28
    ws.column_dimensions["G"].width = 20


# ─── Main ────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="Export Firestore → Excel")
    parser.add_argument("--output", default=None, help="Chemin du fichier de sortie")
    parser.add_argument("--since",  default=None, help="Date minimale YYYY-MM-DD")
    args = parser.parse_args()

    output = args.output or f"transactions-export-{date.today()}.xlsx"

    print("🔑  Connexion à Firestore...")
    db = _init_firebase()

    txs = fetch_transactions(db, since=args.since)
    if not txs:
        print("⚠️  Aucune transaction trouvée.")
        return

    print(f"📊  Construction du classeur Excel ({len(txs)} transactions)...", end=" ", flush=True)
    wb = build_workbook(txs)
    add_stats_sheet(wb, txs)
    print("OK")

    wb.save(output)
    size_kb = Path(output).stat().st_size // 1024
    print(f"✅  Fichier exporté : {output}  ({size_kb} KB, {len(txs)} lignes)")


if __name__ == "__main__":
    main()
